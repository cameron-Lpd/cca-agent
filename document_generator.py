"""
Generates Word (.docx) documents from structured campaign and DM text content.
"""

import re
import base64
from io import BytesIO

from docx import Document
from docx.shared import Pt, RGBColor, Cm


def _base_doc() -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    return doc


def _add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    p2 = doc.add_paragraph()
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()  # spacer


def _add_block(doc: Document, header: str, body: str) -> None:
    # Coloured header
    h = doc.add_paragraph()
    hr = h.add_run(header)
    hr.bold = True
    hr.font.size = Pt(12)
    hr.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    # Body — one paragraph per blank-line-separated chunk
    for chunk in re.split(r'\n{2,}', body.strip()):
        chunk = chunk.strip()
        if not chunk:
            continue
        p = doc.add_paragraph()
        r = p.add_run(chunk)
        r.font.size = Pt(11)
        p.paragraph_format.space_after = Pt(3)

    doc.add_paragraph()  # spacer between blocks


def generate_campaign_docx(name: str, caption: str, detail) -> bytes:
    doc = _base_doc()
    _add_title(
        doc,
        name,
        f"{detail.title}  ·  {detail.date}  ·  {detail.primary_client or ''}",
    )

    # Split on POST N WC headers (keep header on each block)
    blocks = re.split(r'(?m)(?=^\s*POST \d+\b)', caption.strip())
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        header = lines[0].strip()
        body = "\n".join(lines[1:]).strip()
        _add_block(doc, header, body)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_dm_docx(name: str, dm_text: str, detail) -> bytes:
    doc = _base_doc()
    _add_title(doc, name, f"{detail.title}  ·  {detail.date}")

    # Split on MESSAGE N headers
    blocks = re.split(r'(?m)(?=^\s*MESSAGE \d+\b)', dm_text.strip())
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        header = lines[0].strip()
        body = "\n".join(lines[1:]).strip()
        _add_block(doc, header, body)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def to_base64(doc_bytes: bytes) -> str:
    return base64.b64encode(doc_bytes).decode()
